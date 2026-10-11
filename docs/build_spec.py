#!/usr/bin/env python3
# Ringató Virággal: reprodukálható képernyőképes HTML + DOCX specifikáció
from pathlib import Path
import threading, functools, http.server, html, datetime
from playwright.sync_api import sync_playwright
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from PIL import Image
ROOT=Path(__file__).resolve().parent.parent
DOC=ROOT/'docs';IMG=DOC/'screens';IMG.mkdir(parents=True,exist_ok=True)
DATE='2026. október 11.';VER='V1.2'
F=[
('F01','Belépés, profil','WP felhasználói fiók, saját adatok, jogosultságok. A demóban Anna fiktív fiókja látszik.'),
('F02','Kezdőlap','Ajánlott médiák, közelgő programok, bérlet és előfizetés státusz.'),
('F03','Médiatár','Kategóriák, keresés, ingyenes és előfizetéses hang-/videótartalmak.'),
('F04','Lejátszó','Fekvő és álló videó, lejátszás, kedvenc, sor, teljes képernyő, tájolásváltás.'),
('F05','Kedvencek / sor','Médiatartalmak mentése és lejátszási sorrend; fiókhoz kötött tárolás.'),
('F06','App-előfizetés','Egy csomag: 8 000 Ft/hó, automatikusan megújuló bankkártyás terhelés.'),
('F07','Előfizetés lemondás','Következő levonás kikapcsolható; a már fizetett időszak végéig használható az app.'),
('F08','Termékkatalógus','Összes aktívan megvásárolható WooCommerce-termék listája, webes ár/készlet alapján.'),
('F09','Termékvásárlás','WooCommerce kosár/rendelés, Barion fizetés, megerősítés utáni jogosultság.'),
('F10','Korábbi rendelések','A felhasználó saját vásárlásainak listája, dátum, összeg, rendelés státusza.'),
('F11','Bérletek','Fennmaradó/összes alkalmak, érvényesség, QR és felhasználási események.'),
('F12','Rendszeres program','Nincs létszámlimit, nincs előzetes foglalás; jegy/bérlet lehet szükséges.'),
('F13','Limitált ingyenes program','Előzetes foglalás és lemondás, a kapacitás szerveroldali ellenőrzésével.'),
('F14','Limitált fizetős program','Eseményhez rendelt külön WooCommerce-termék; sikeres fizetés = helyfoglalás.'),
('F15','Beléptetés','Virág/staff QR ellenőrzés, bérletalkalom egyszeri, naplózott levonása.'),
('F16','Üzenetek','Szülő és Virág közti aszinkron szöveges üzenetküldés.'),
('F17','Értesítések','Új média/program és foglalási értesítések, PWA web push engedéllyel.'),
('F18','Kérdések','Rövid szavazások, egyszeri válasz és egyszerű eredmények.'),
('F19','Admin','Meglévő WP/WooCommerce admin; a PWA-ban csak mobilos staff beléptetés.'),
('F20','Megszűnő funkciók','Nincs Kuponjaim, rövid idejű olcsó médiahozzáférés vagy új Online Ringató értékesítés.')
]
E=[
('WP User','id, név, e-mail, jogosultsági szerep','Rendelések, bérletek, foglalások, üzenetek'),
('WC Product','id, cím, ár, terméktípus, aktív, megvehető, készlet','Rendeléstétel; opcionálisan esemény/előfizetés'),
('WC Order / Item','order_id, tételek, összeg, státusz, user_id, dátum','Fizetési eredményből jegy/bérlet-jogosultság'),
('Subscription Plan','product_id, 8000 Ft, 1 hónap, recurring','WC subscription termék'),
('Subscription','id, user_id, státusz, következő terhelés, fizetett időszak vége','Előfizetési terv + fizetési események'),
('Payment Transaction','provider_id, order_id, összeg, státusz, idempotency_key','Barion ↔ WooCommerce; ellenőrzött webhook'),
('Event','id, név, hely, kezdet, típus, kapacitás?, product_id?','Foglalások, eseményjegyek'),
('Reservation','id, event_id, user_id, állapot, létszám, order_id?','Limitált program + résztvevő; fizetett rendelés'),
('Ticket','id, order_item_id, owner_id, event_id?, QR token, státusz','Vásárolt jegy, esetleg eseményhely'),
('Pass','id, owner_id, total, remaining, valid_until','Bérletfelhasználás, rendelési tétel'),
('Pass Usage','id, pass_id, staff_id, event_id?, timestamp, unique_key','Beléptetés, duplikált levonás kizárása'),
('Media','id, cím, típus, képarány, hossz, kategória, premium, cdn_ref','Kategóriák, kedvencek, lejátszási sor'),
('Media Category','id, név, sorrend, aktív','N:M média'),
('Favorite / Queue','user_id, media_id, position, created_at','N:M user ↔ média'),
('Message','id, thread_id, sender_id, receiver_id, text, read_at','Szülő ↔ Virág beszélgetés'),
('Poll / Response','poll_id, kérdés, opciók; response_id, user_id, választás','Szavazásonként egy válasz felhasználónként'),
('Push Subscription','id, user_id, endpoint, keys, permission, updated_at','Eszköz / böngésző ↔ felhasználó')
]
ST=[
('Előfizetés','inaktív → fizetés → aktív → lemondott (aktív a periódus végéig) → lejárt'),
('Alapfoglalkozás','látható → látogatható; nincs foglalás vagy férőhely'),
('Ingyenes limitált','szabad hely → foglalás → megerősítve / lemondva / betelt'),
('Fizetős limitált','termék → fizetés függőben → fizetett jegy + foglalás / sikertelen'),
('Bérlet','aktív → alkalmak csökkennek → elfogyott vagy lejárt'),
('Média','ingyenes / zárolt / betölt / lejátszás / szünet / hiba / teljes képernyő'),
('Push','nem engedélyezett / engedélyezve / letiltott')
]
FLOWS=[
('App havi előfizetés','Vásárlás → havi 8 000 Ft termék → WooCommerce Subscriptions → Barion recurring engedély → szerveroldali visszaigazolás → teljes médiatár. Lemondás → következő terhelés leáll → hozzáférés időszak végéig.'),
('Rendszeres foglalkozás','Programok → alapfoglalkozás → részletek. Nincs foglalásgomb, nincs létszámlimit. Jegy vagy bérlet vásárolható.'),
('Limitált, ingyenes','Programok → szabad helyek → helyfoglalás → visszaigazolás → Foglalásaim; lemondás felszabadítja a helyet.'),
('Limitált, fizetős','Programok → WooCommerce eseményjegy → Barion fizetés → fizetett rendelés → automatikusan létrejött jegy/foglalás. Sikertelen fizetés nem foglal végleges helyet.'),
('Bérlet és beléptetés','Jegyeim → QR → Virág staff jogosultsággal beolvassa → validáció → egyetlen naplózott alkalomlevonás.'),
('Médiatár','Kezdőlap / Tár → kategória / keresés → tartalom → ingyenes vagy fizetett jogosultság → lejátszás; különben előfizetési ajánlat.')
]
VIDEO=[
('16:9 videó, álló telefon','A teljes fekvő képkocka látható legyen; levágás/crop nélkül, oldalsó vagy felső üres sáv megengedett.'),
('9:16 videó, álló telefon','Az álló videó maximális méretben, biztonságos vezérlőkkel és feliratokkal látszódjon.'),
('16:9 videó, fekvő telefon','Nagyméretű, lehetőleg fullscreen kép; safe-area és navigáció kontroll.'),
('9:16 videó, fekvő telefon','Teljes álló tartalom középre illesztve, oldalt kitöltő sávok; tilos crop-pal széthúzni.'),
('Orientációváltás','Forgatáskor az időpozíció megmarad; responsive layout és fullscreen visszatérés működik.'),
('Kompatibilitás','iPhone Safari telepített PWA és Android Chrome PWA, használható MP4/HLS, playsinline, object-fit:contain.'),
('Jogosultság és minőség','Védett, rövid életű stream URL, thumbnail, buffer/hiba/újrapróbálás; felirat szükség esetén.')
]
SC=[
('01-home','Főoldal / ingyenes hozzáférés','home','', 'Ajánlott dalok, soron következő programok, bérlet és előfizetés jelzése.','Főoldal → Tár / Programok / Vásárlás / Jegyeim.'),
('02-library','Médiatár / zárolt elemek','library','', 'Keresés, kategóriák és prémium lakat; ingyenes tartalom közvetlenül indul.','Tár → tartalom → lejátszó vagy előfizetési oldal.'),
('03-player-wide','Fekvő videó lejátszó','library','openContent("hinta")', '16:9 képarány; teljes képkocka, vezérlők és kedvencek.','Tár → média → lejátszó → kedvenc / sor.'),
('04-paywall','Zárolt tartalom / előfizetési ajánlat','library','openContent("csip")', 'Havi 8 000 Ft-os csomag; nincs olcsó rövid hozzáférés.','Zárolt tartalom → Vásárlás → előfizetés.'),
('05-library-active','Médiatár / aktív előfizetés','library','ACTIVE', 'A prémium tartalom a jogosultság megszerzése után elérhető.','Előfizetéses vásárlás → médiatár.'),
('06-player-vertical','Álló videó lejátszó','library','ACTIVE;openContent("mesel")', '9:16 képarány; álló és fekvő telefonon is teljes tartalom.','Tár → függőleges videó → teljes képernyő.'),
('07-events','Programlista / három típus','calendar','', 'Rendszeres, ingyenes foglalós és fizetős foglalós program.','Programok → a megfelelő részletfolyamat.'),
('08-book-free','Ingyenes programfoglalás','calendar','eventDetails("limited")', 'A férőhely-limit miatt foglalás szükséges.','Program → helyfoglalás → Foglalásaim / lemondás.'),
('09-book-paid','Fizetős esemény / jegy','calendar','eventDetails("concert")', 'Sikeres jegyvásárlás egyben foglalás.','Program → WooCommerce-jegy → Barion → foglalás.'),
('10-shop','Vásárlás / induló állapot','store','', 'Aktív termékek, új app havi előfizetés; korábbi Online Ringató nincs kínálva.','Vásárlás → termék / előfizetés / rendelési előzmények.'),
('11-sub-active','Havi előfizetés / aktív','store','ACTIVE', 'Automatikusan megújuló havi tagság, 8 000 Ft.','Vásárlás → lemondási lehetőség.'),
('12-sub-ending','Havi előfizetés / lemondva','store','ACTIVE;ENDING', 'A következő terhelés leáll, a fizetett időszak végéig hozzáfér.','Lemondás → státusz, lejárat → opcionális újraaktiválás.'),
('13-wallet','Jegyek, bérlet, korábbi rendelések','passes','', '8 alkalmas bérlet maradéka, vásárolt jegyek, rendelési előzmények.','Vásárlás → Jegyeim → QR.'),
('14-qr','Bérlet / QR-bemutatás','passes','showQR()', 'Bemutatáskor nincs levonás; csak a sikeres helyszíni ellenőrzés számít.','Jegyeim → QR → beléptetés.'),
('15-checkin','Virág / beléptetés','checkin','scan()', 'Staff QR ellenőrzés; alkalomlevonás csak sikeres beváltáskor.','Jogosult staff → beolvasás → naplózott használat.'),
('16-messages','Üzenetek','messages','', 'Szülő és Virág aszinkron üzenetei.','Alsó Üzenet lap → beszélgetés.'),
('17-profile','Profil / állapotok','home','showProfile()', 'Előfizetés kezelése, értesítések és demo reset.','Fejléc profil → előfizetés menü.'),
('18-notifications','Értesítések','home','showNotifications()', 'Új tartalom- és programértesítések.','Fejléc csengő → kapcsolódó tartalom/program.')
]
QUESTIONS=[
('WooCommerce katalogizálás','Az összes aktuális termék és kategória listája, archivált termékek szűrése.'),
('Előfizetés','Név, havi számlázási nap, ÁSZF, fogyasztói tájékoztatás, számlázási folyamat.'),
('Barion','Recurring/token engedély, WooCommerce Subscriptions kompatibilis bővítmény.'),
('Online Ringató kivezetése','Meglévő vásárlók jogosultságai, archiválás, átmenet.'),
('Foglalási szabályok','Család vagy személy számít kapacitásnak? Várólista és lemondási határidő?'),
('Fizetős események','Átmeneti készletfoglalás, visszatérítés, lemondás.'),
('Bérletek','Érvényesség, használat többször egy napon, átruházás, foglalkozás-típus jogosultság.'),
('Média','CDN, videóminőség, képformátum, feliratozás, tartalomhozzáférés.'),
('Adatkezelés','Értesítések, személyes üzenetek, hozzájárulások, megőrzés.')
]
class Server(http.server.SimpleHTTPRequestHandler):
    def log_message(self,*args): pass
def capture():
    serve=http.server.ThreadingHTTPServer(('127.0.0.1',8826),functools.partial(Server,directory=str(ROOT)))
    threading.Thread(target=serve.serve_forever,daemon=True).start()
    try:
        with sync_playwright() as p:
            br=p.chromium.launch(headless=True,args=['--no-sandbox'])
            page=br.new_page(viewport={'width':414,'height':900},device_scale_factor=1)
            page.goto('http://127.0.0.1:8826/app.html')
            for filename,title,screen,setup,description,nav in SC:
                page.evaluate("() => {localStorage.removeItem('ringato-v1'); location.reload();}")
                page.wait_for_load_state('load')
                page.wait_for_timeout(130)
                if 'ACTIVE' in setup:
                    page.evaluate("() => {state.subscription={status:'active',renew:true,next:'2026.11.11.'};state.plus=true;save();}")
                if 'ENDING' in setup:
                    page.evaluate("() => {state.subscription.status='ending';state.subscription.renew=false;save();}")
                page.evaluate("(s) => go(s)",screen)
                cmd=setup.replace('ACTIVE','').replace('ENDING','').strip(' ;')
                if cmd:page.evaluate("() => {"+cmd+";}")
                page.wait_for_timeout(120)
                page.screenshot(path=str(IMG/(filename+'.png')),full_page=True,animations='disabled')
            br.close()
    finally:serve.shutdown()

CSS="""*{box-sizing:border-box}body{margin:0;background:#fbf4e8;color:#33251f;font-family:system-ui,-apple-system,sans-serif;line-height:1.55}
header{background:linear-gradient(135deg,#a90012,#7d0010);color:white;padding:43px max(24px,calc((100vw - 1100px)/2));border-radius:0 0 65px 0}
header h1{font:700 39px Georgia,serif;margin:10px 0}header p{color:#fff6ee}main{max-width:1100px;margin:auto;padding:25px}
nav{display:flex;flex-wrap:wrap;gap:12px;background:white;border-radius:18px;padding:18px;margin-bottom:24px}a{color:#a90012}
section{background:white;margin:25px 0;padding:27px;border:1px solid #eadccf;border-radius:22px 22px 40px 22px}
h2{font:700 26px Georgia,serif;color:#a90012;margin:0 0 17px}h3{font:700 21px Georgia,serif;color:#750010}
table{width:100%;border-collapse:collapse;font-size:13px}th,td{border-bottom:1px solid #e8ddd4;vertical-align:top;text-align:left;padding:10px}th{background:#fce8b3}
figure{display:grid;grid-template-columns:260px 1fr;gap:25px;margin:0;padding:25px 0;border-bottom:1px solid #e9ded0}
figure img{width:100%;max-height:700px;object-fit:contain;object-position:top;border:6px solid #efd9c5;border-radius:26px}
figcaption h3{margin:0 0 8px}.navline{background:#fff2d1;border-left:4px solid #efb229;padding:12px}
.note{background:#fff0dd;border-radius:13px;padding:16px}.badge{background:#f6cb5e;padding:7px 12px;border-radius:30px;font-weight:800;color:#5b2515;font-size:12px}
@media(max-width:720px){header h1{font-size:30px}main{padding:10px}section{padding:17px}figure{grid-template-columns:1fr}figure img{max-width:280px;margin:auto}}
"""
def tbl(headers,rows):
    return '<div style="overflow-x:auto"><table><tr>'+''.join('<th>'+html.escape(h)+'</th>' for h in headers)+'</tr>'+''.join('<tr>'+''.join('<td>'+html.escape(str(v))+'</td>' for v in row)+'</tr>' for row in rows)+'</table></div>'
def make_html():
    b=['<!doctype html><html lang="hu"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>Ringató PWA · V1.2 specifikáció</title><style>'+CSS+'</style></head><body>']
    b+=['<header><span class="badge">EGYEZTETÉSI SPECIFIKÁCIÓ · V1.2</span><h1>Ringató Virággal<br>Mobil PWA</h1><p>'+DATE+' · funkciók, adatmodell, folyamatok és állapotfüggő képernyőképek</p><p><a style="color:white" href="../app.html">Interaktív prototípus megnyitása →</a></p></header>']
    anchors=[('overview','Cél és hatókör'),('features','Funkciók'),('entities','Entitások'),('states','Állapotok'),('flows','Folyamatok'),('video','Videó'),('screens','Képernyők'),('integration','Integráció'),('decisions','Nyitott kérdések')]
    b+=['<main><nav>'+''.join('<a href="#'+x+'">'+y+'</a>' for x,y in anchors)+'</nav>']
    b+=['<section id="overview"><h2>01 · Cél és hatókör</h2><p>Az alkalmazás a meglévő WordPress/WooCommerce weboldalhoz kapcsolódó, mobilra telepíthető PWA. <b>Nem külön webshop és nem külön adminrendszer.</b></p><p>Vásárláskor az összes aktuálisan megvásárolható WooCommerce-terméket listázni kell, továbbá a külön <b>8 000 Ft/hó Ringató App előfizetést</b>, automatikus megújulással és lemondással. A korábbi Online Ringató kivezetésre kerül. Nincs Kuponjaim funkció, nincs olcsóbb rövid hozzáférés.</p><p class="note"><b>A jelenlegi GitHub Pages demó</b> nem végez valós fizetést, foglalást, bérletlevonást vagy üzenetküldést; a terméklista mintakínálat. Az éles rendszerben WordPress/WooCommerce a kizárólagos adatforrás.</p></section>']
    b+=['<section id="features"><h2>02 · Funkciólista</h2>'+tbl(['ID','Funkció','Leírás'],F)+'</section>']
    b+=['<section id="entities"><h2>03 · Kezelt entitások és kapcsolatok</h2>'+tbl(['Entitás','Fő tulajdonságok','Kapcsolat'],E)+'</section>']
    b+=['<section id="states"><h2>04 · Állapotok</h2>'+tbl(['Terület','Állapotgép'],ST)+'</section>']
    b+=['<section id="flows"><h2>05 · Üzleti folyamatok és navigáció</h2>'+''.join('<h3>'+html.escape(x)+'</h3><p>'+html.escape(y)+'</p>' for x,y in FLOWS)+'</section>']
    b+=['<section id="video"><h2>06 · Álló és fekvő videó – kötelező</h2><p>Vegyes <b>16:9 fekvő</b> és <b>9:16 álló</b> videók, amelyek <b>álló és fekvő telefonhelyzetben</b> is torzítás és képkivágás nélkül lejátszhatók. Mind a négy kombinációt külön kell tesztelni. A böngésző tájolásváltásakor a lejátszás folytatódik, a kezelőgombok láthatók maradnak.</p>'+tbl(['Média × telefon','Követelmény'],VIDEO)+'<p><b>Műszaki alapelv:</b> eredeti képarány, object-fit: contain, playsinline, megfelelő fullscreen, mobil-kompatibilis streaming, védett média-URL és jól látható vezérlők.</p></section>']
    b+=['<section id="screens"><h2>07 · Képernyők – több állapotban</h2><p>A képernyőképek az aktuális prototípus tényleges böngésző-rendereléséből készülnek. A funkciók demóadatokon működnek.</p>']
    for i,(name,title,screen,setup,desc,nav) in enumerate(SC,1):
        b+=['<figure id="'+name+'"><img loading="lazy" src="screens/'+name+'.png" alt="'+html.escape(title)+'"><figcaption><p style="font-size:12px;color:#7e6652">KÉPERNYŐ '+str(i).zfill(2)+'</p><h3>'+html.escape(title)+'</h3><p>'+html.escape(desc)+'</p><p class="navline"><b>Navigáció és összefüggés:</b><br>'+html.escape(nav)+'</p><p style="font-size:12px">Prototípus állapot: '+html.escape(setup or screen)+'</p></figcaption></figure>']
    b+=['</section>']
    b+=['<section id="integration"><h2>08 · Architekturális és nem-funkcionális követelmények</h2><h3>Integráció</h3><p>WordPress-fiók és admin; WooCommerce-terméktörzs, rendelés, előfizetés és fizetés; Barion recurring engedélyezett integrációval. Ringató WP plugin REST API: program, foglalás, bérlet, jegy, jogosultság, média, üzenet, push. PWA ajánlottan ugyanazon domain /app útvonalán. Média biztonságos CDN-ről.</p><h3>Tranzakciók és jogok</h3><p>Webhook idempotencia, sikertelen fizetés esetén nem készül végleges foglalás. Kapacitást szerveroldali zárolás védi. QR-ben nincs személyes adat; staff jogosultság és auditnapló. Az előfizetés lejáratát szerver ellenőrzi.</p><h3>Minőség és biztonság</h3><p>HTTPS, auth + capability check, CSRF, GDPR adatminimalizálás, tartalomvédelem, mentések, staging környezet, regressziós WooCommerce/Barion tesztek. iOS/Android PWA, négy videóorientáció, fizetési hiba, kapcsolatszakadás és több eszköz tesztelendő.</p></section>']
    b+=['<section id="decisions"><h2>09 · Nyitott döntések</h2>'+tbl(['Téma','Egyeztetés'],QUESTIONS)+'</section></main><footer style="max-width:1100px;margin:30px auto;padding:20px;color:#786c64">Ringató Virággal · V1.2 · egyeztetési tervezet, nem éles integráció · '+DATE+'</footer></body></html>']
    (DOC/'specification.html').write_text(''.join(b),encoding='utf-8')
def make_docx():
    d=Document();sec=d.sections[0];sec.top_margin=Inches(.75);sec.bottom_margin=Inches(.72);sec.left_margin=Inches(.75);sec.right_margin=Inches(.75)
    d.styles['Normal'].font.size=Pt(9);d.styles['Normal'].font.name='Aptos'
    for x,size in [('Title',27),('Heading 1',17),('Heading 2',12)]:
        st=d.styles[x];st.font.name='Aptos Display';st.font.size=Pt(size);st.font.color.rgb=RGBColor(169,0,18)
    h=sec.header.paragraphs[0];h.text='RINGATÓ VIRÁGGAL / V1.2 FUNKCIONÁLIS SPECIFIKÁCIÓ';h.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    d.add_paragraph('Ringató Virággal – Mobil PWA',style='Title');d.add_paragraph('Funkcionális specifikáció · '+DATE)
    d.add_heading('1. Cél, scope és termékszabályok',1)
    d.add_paragraph('Meglévő WordPress/WooCommerce rendszer kiegészítése; az app nem külön webshop vagy teljes admin. Az összes aktív megvásárolható termék listázása éles WooCommerce-adatokból; külön 8 000 Ft/hó automatikusan megújuló Ringató App-előfizetés, mely lemondható. Az Online Ringató kivezetésre kerül. Nincs Kuponjaim funkció és nincs rövid idejű hozzáférés.')
    d.add_paragraph('A GitHub Pages demó statikus, fizetés és foglalás nem valós. A képernyőképek a jelenlegi prototípusból készültek, mintadatokkal.')
    def table(headers,rows):
        t=d.add_table(rows=1,cols=len(headers));t.style='Light Shading Accent 1'
        for i,x in enumerate(headers):t.rows[0].cells[i].text=x
        for row in rows:
            rr=t.add_row()
            for i,x in enumerate(row):rr.cells[i].text=str(x)
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.paragraph_format.space_after=Pt(2)
                    for r in p.runs:r.font.size=Pt(8)
        d.add_paragraph()
    d.add_heading('2. Funkciólista',1);table(['ID','Funkció','Rövid magyarázat'],F)
    d.add_heading('3. Entitások és tulajdonságok',1);table(['Entitás','Tulajdonságok','Kapcsolatok'],E)
    d.add_heading('4. Állapotok',1);table(['Terület','Állapotok'],ST)
    d.add_heading('5. Navigáció és folyamatok',1)
    for x,y in FLOWS:d.add_heading(x,2);d.add_paragraph(y)
    d.add_heading('6. Videólejátszás: álló és fekvő',1)
    d.add_paragraph('Kötelező minden 16:9 és 9:16 videó teljes képkockás, torzításmentes lejátszása álló és fekvő mobilon. A négy kombináció külön elfogadási teszt. Orientációváltás megőrzi a pozíciót, fullscreen támogatott.')
    table(['Kombináció / téma','Követelmény'],VIDEO)
    d.add_heading('7. Képernyőképek, állapotok és navigáció',1)
    for i,(name,title,screen,setup,desc,nav) in enumerate(SC,1):
        d.add_heading('7.'+str(i)+' '+title,2)
        t=d.add_table(rows=1,cols=2);t.autofit=False;t.columns[0].width=Inches(2.35);t.columns[1].width=Inches(3.85)
        shot=IMG/(name+'.png')
        with Image.open(shot) as im: w,h=im.size
        width_in=min(2.10,6.12*w/h)
        t.cell(0,0).paragraphs[0].add_run().add_picture(str(shot),width=Inches(width_in))
        t.cell(0,1).text='Funkció / állapot:\n'+desc+'\n\nNavigáció:\n'+nav+'\n\nPrototípus: '+(setup or screen)
        trPr=t.rows[0]._tr.get_or_add_trPr();trPr.append(OxmlElement('w:cantSplit'))
        d.add_paragraph()
    d.add_heading('8. Architekturális, biztonsági és elfogadási elvek',1)
    d.add_paragraph('WordPress/WooCommerce egyetlen termék- és rendelésforrás. Ringató WP REST plugin kezeli az egyedi eseményeket, foglalásokat, jegyeket, bérleteket és médiákat. Barion recurring jogosultságot ellenőrizni kell. Szerveroldali idempotens webhookok, kapacitászár, staff-belépés és audit QR-levonás. Privát média-CDN és hitelesített hozzáférés. HTTPS, GDPR, staging, rollback és regressziós tesztek szükségesek.')
    d.add_heading('9. Nyitott kérdések',1);table(['Téma','Egyeztetendő'],QUESTIONS)
    out=DOC/'Ringato_V1_2_specifikacio.docx';d.save(out)
    print('Created',out, out.stat().st_size)
if __name__=='__main__':
    capture();make_html();make_docx()
